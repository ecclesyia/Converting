import os
import fitz
from PIL import Image
import io
from docx import Document
import openpyxl
from pptx import Presentation
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

def images_to_pdf(image_paths: list, pdf_path: str):
    """Combines one or multiple images (PNG, JPG, WEBP, BMP, etc.) into a unified PDF."""
    doc = fitz.open()
    for img_path in image_paths:
        img = Image.open(img_path)
        if img.mode in ('RGBA', 'LA') or (img.mode == 'P' and 'transparency' in img.info):
            img = img.convert('RGB')
        
        img_bytes_io = io.BytesIO()
        img.save(img_bytes_io, format='JPEG', quality=95)
        img_bytes = img_bytes_io.getvalue()
        
        img_doc = fitz.open("pdf", fitz.open("jpeg", img_bytes).convert_to_pdf())
        doc.insert_pdf(img_doc)
    
    doc.save(pdf_path)
    doc.close()
    return pdf_path

def text_to_pdf(text_or_path: str, pdf_path: str, is_file: bool = True):
    """Converts plain text or markdown content into a neatly formatted PDF."""
    if is_file:
        with open(text_or_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
    else:
        content = text_or_path

    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            rightMargin=40, leftMargin=40,
                            topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    normal_style = styles['Normal']
    normal_style.fontSize = 10
    normal_style.leading = 14

    story = []
    lines = content.split('\n')
    for line in lines:
        cleaned = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        if not cleaned.strip():
            story.append(Spacer(1, 8))
        else:
            story.append(Paragraph(cleaned, normal_style))
    
    doc.build(story)
    return pdf_path

def docx_to_pdf(docx_path: str, pdf_path: str):
    """Converts DOCX Word document to PDF."""
    try:
        import comtypes.client
        import pythoncom
        pythoncom.CoInitialize()
        word = comtypes.client.CreateObject('Word.Application')
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        doc.Close()
        word.Quit()
        return pdf_path
    except Exception:
        pass

    doc = Document(docx_path)
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                                rightMargin=40, leftMargin=40,
                                topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            story.append(Spacer(1, 6))
            continue
        cleaned = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        if para.style.name.startswith('Heading 1'):
            p_style = styles['Heading1']
        elif para.style.name.startswith('Heading 2'):
            p_style = styles['Heading2']
        elif para.style.name.startswith('Heading 3'):
            p_style = styles['Heading3']
        elif para.style.name.startswith('Title'):
            p_style = styles['Title']
        else:
            p_style = styles['Normal']
        
        story.append(Paragraph(cleaned, p_style))
        story.append(Spacer(1, 4))

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            t = Table(table_data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F3F4F6')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#111827')),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
                ('FONTSIZE', (0,0), (-1,-1), 9),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D1D5DB')),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    if not story:
        story.append(Paragraph("Empty Document", styles['Normal']))

    pdf_doc.build(story)
    return pdf_path

def excel_to_pdf(xlsx_path: str, pdf_path: str):
    """Converts Excel workbook worksheets to formatted table pages in PDF."""
    try:
        import comtypes.client
        import pythoncom
        pythoncom.CoInitialize()
        excel = comtypes.client.CreateObject('Excel.Application')
        excel.Visible = False
        wb = excel.Workbooks.Open(os.path.abspath(xlsx_path))
        wb.ExportAsFixedFormat(0, os.path.abspath(pdf_path))
        wb.Close(False)
        excel.Quit()
        return pdf_path
    except Exception:
        pass

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                                rightMargin=30, leftMargin=30,
                                topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    story = []

    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        story.append(Paragraph(f"<b>Sheet: {sheetname}</b>", styles['Heading2']))
        story.append(Spacer(1, 6))

        data = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                formatted_row = [str(c) if c is not None else "" for c in row]
                data.append(formatted_row)
        
        if data:
            max_cols = max(len(r) for r in data)
            normalized = [r + [""]*(max_cols - len(r)) for r in data]
            t = Table(normalized)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E5E7EB')),
                ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
                ('FONTSIZE', (0,0), (-1,-1), 8),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#9CA3AF')),
                ('TOPPADDING', (0,0), (-1,-1), 3),
                ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 14))

    if not story:
        story.append(Paragraph("Empty Workbook", styles['Normal']))

    pdf_doc.build(story)
    return pdf_path

def pptx_to_pdf(pptx_path: str, pdf_path: str):
    """Converts PowerPoint presentation to PDF."""
    try:
        import comtypes.client
        import pythoncom
        pythoncom.CoInitialize()
        powerpoint = comtypes.client.CreateObject('PowerPoint.Application')
        pres = powerpoint.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)
        pres.SaveAs(os.path.abspath(pdf_path), 32)
        pres.Close()
        powerpoint.Quit()
        return pdf_path
    except Exception:
        pass

    prs = Presentation(pptx_path)
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=letter,
                                rightMargin=40, leftMargin=40,
                                topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    for i, slide in enumerate(prs.slides):
        story.append(Paragraph(f"<b>Slide {i+1}</b>", styles['Heading2']))
        story.append(Spacer(1, 6))

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    txt = paragraph.text.strip()
                    if txt:
                        cleaned = txt.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        story.append(Paragraph(cleaned, styles['Normal']))
                        story.append(Spacer(1, 4))
        story.append(Spacer(1, 12))

    if not story:
        story.append(Paragraph("Empty Presentation", styles['Normal']))

    pdf_doc.build(story)
    return pdf_path
